"""DOCX parsing strategy — text, tables, and embedded images."""
import logging
from pathlib import Path

from .base import ParserStrategy

logger = logging.getLogger(__name__)


class DocxParserStrategy(ParserStrategy):
    """Extracts paragraphs, tables (as Markdown), and describes embedded images."""

    def parse(self, file_path: Path) -> str:
        import pandas as pd
        from docx import Document as DocxDocument

        logger.info(f"Parsing DOCX: {file_path}")
        try:
            doc = DocxDocument(file_path)
            content = []

            for para in doc.paragraphs:
                content.append(para.text)

            for i, table in enumerate(doc.tables):
                rows = []
                for row in table.rows:
                    row_data = []
                    for cell in row.cells:
                        text = cell.text
                        if "w:drawing" in cell._element.xml or "v:imagedata" in cell._element.xml:
                            text = f"{text} IMAGE_FOUND".strip()
                        row_data.append(text)
                    rows.append(row_data)
                if rows:
                    df = pd.DataFrame(rows)
                    content.append(f"\n[Table {i + 1}]:\n{df.to_markdown(index=False, headers='keys')}\n")

            processed = set()
            image_count = 0
            for rel in doc.part.rels.values():
                if not (hasattr(rel, "target_ref") and "image" in rel.target_ref):
                    continue
                if not hasattr(rel, "target_part"):
                    continue
                blob = rel.target_part.blob
                if blob in processed:
                    continue
                processed.add(blob)
                image_count += 1
                try:
                    desc = self.image_describer.describe(blob, context=f"Image {image_count} in document")
                    content.append(desc)
                except Exception as e:
                    logger.warning(f"Image {image_count} description failed: {e}")

            return "\n".join(content)
        except Exception as e:
            logger.error(f"DOCX parsing failed: {e}")
            return ""
